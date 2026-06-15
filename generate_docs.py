#!/usr/bin/env python3
"""Генерация пояснительной записки и отзыва для курсовой работы."""

from docx import Document
from docx.shared import Pt, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from fpdf import FPDF
import os


def set_cell_border(cell, **kwargs):
    """Установить границы ячейки таблицы."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        if edge in kwargs:
            edge_data = kwargs.get(edge)
            if edge_data:
                tag = f'w:{edge}'
                element = tcPr.find(tag)
                if element is None:
                    element = parse_xml(r'<w:%s xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>' % edge)
                    tcPr.append(element)
                element.set(qn('w:val'), 'single')
                element.set(qn('w:sz'), '4')
                element.set(qn('w:space'), '0')
                element.set(qn('w:color'), '000000')


from docx.oxml import parse_xml
from docx.oxml.ns import qn


def create_explanatory_note_docx(filename):
    doc = Document()

    # Настройки страницы
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(1.5)

    # Титульная страница
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("\n\n\n\n")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("МИНИСТЕРСТВО НАУКИ И ВЫСШЕГО ОБРАЗОВАНИЯ\n")
    run.bold = True
    run.font.size = Pt(14)
    run = p.add_run("РОССИЙСКОЙ ФЕДЕРАЦИИ\n\n")
    run.bold = True
    run.font.size = Pt(14)
    run = p.add_run("Федеральное государственное бюджетное образовательное учреждение\n")
    run.font.size = Pt(12)
    run = p.add_run("высшего образования\n")
    run.font.size = Pt(12)
    run = p.add_run("«Название университета»\n")
    run.bold = True
    run.font.size = Pt(14)

    p = doc.add_paragraph("\n\n\n\n")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("ПОЯСНИТЕЛЬНАЯ ЗАПИСКА\n")
    run.bold = True
    run.font.size = Pt(16)
    run = p.add_run("к курсовой работе\n")
    run.font.size = Pt(14)
    run = p.add_run("по дисциплине «Программирование»\n")
    run.font.size = Pt(14)
    run = p.add_run("\nТема: «Разработка веб-приложения для планирования поездок»\n")
    run.bold = True
    run.font.size = Pt(14)

    p = doc.add_paragraph("\n\n\n\n")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("Выполнил: студент ___ курса\n")
    run.font.size = Pt(12)
    run = p.add_run("группы ___\n")
    run.font.size = Pt(12)
    run = p.add_run("И. О. Фамилия\n\n")
    run.font.size = Pt(12)
    run = p.add_run("Проверил: _______________\n")
    run.font.size = Pt(12)

    p = doc.add_paragraph("\n\n\n\n")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("2026\n")
    run.font.size = Pt(12)

    doc.add_page_break()

    # Задание
    heading = doc.add_heading("ЗАДАНИЕ", level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(
        "1. Тема курсовой работы: «Разработка веб-приложения для планирования поездок»."
    )
    doc.add_paragraph(
        "2. Цель работы: разработать полноценное веб-приложение для планирования туристических поездок, "
        "включающее клиентскую и серверную части, систему аутентификации, инструменты для составления маршрута, "
        "учёта расходов и формирования списка вещей."
    )
    doc.add_paragraph(
        "3. Технологический стек: Python, Django, Django REST Framework, React, TypeScript, Docker."
    )
    doc.add_paragraph(
        "4. Требования к результату: работающее веб-приложение, развёрнутое в Docker-контейнерах, "
        "с документацией и инструкцией по развёртыванию."
    )

    doc.add_page_break()

    # Содержание
    doc.add_heading("СОДЕРЖАНИЕ", level=1).alignment = WD_ALIGN_PARAGRAPH.CENTER
    contents = [
        "1. Введение",
        "2. Проектирование приложения",
        "3. Технологический стек",
        "4. Реализация серверной части",
        "5. Реализация клиентской части",
        "6. Тестирование",
        "7. Заключение",
        "Список использованных источников",
    ]
    for item in contents:
        p = doc.add_paragraph(item)
        p.paragraph_format.left_indent = Cm(1)

    doc.add_page_break()

    # Основной текст
    doc.add_heading("1. Введение", level=1)
    doc.add_paragraph(
        "В современном мире путешествия стали неотъемлемой частью жизни. Для эффективной организации "
        "поездок необходимо учитывать множество факторов: маршрут, расходы, даты, список необходимых вещей. "
        "Ручное планирование часто приводит к ошибкам и недопониманию. "
        "Веб-приложение для планирования поездок позволяет систематизировать весь процесс подготовки к путешествию."
    )

    doc.add_heading("2. Проектирование приложения", level=1)
    doc.add_paragraph(
        "Приложение построено по клиент-серверной архитектуре. Серверная часть реализована на Django "
        "с использованием Django REST Framework и предоставляет API для работы с данными. "
        "Клиентская часть разработана на React с использованием TypeScript и взаимодействует с сервером по REST API."
    )
    doc.add_paragraph(
        "Основные сущности приложения: пользователь, поездка, пункт маршрута, расход, вещь в чемодане, "
        "бронирование и настройки пользователя."
    )

    doc.add_heading("3. Технологический стек", level=1)
    tech = [
        "Backend: Python 3.10, Django 5, Django REST Framework, JWT-аутентификация",
        "Frontend: React 18, TypeScript, Vite, Tailwind CSS",
        "База данных: SQLite (для разработки)",
        "Контейнеризация: Docker, Docker Compose, Nginx",
        "Контроль версий: Git, GitHub",
    ]
    for item in tech:
        p = doc.add_paragraph(item, style='List Bullet')

    doc.add_heading("4. Реализация серверной части", level=1)
    doc.add_paragraph(
        "Серверная часть включает приложения accounts (аутентификация и настройки пользователя) и "
        "trips (управление поездками, маршрутом, расходами, бронированиями). "
        "API построено на основе ViewSet'ов Django REST Framework. Аутентификация выполняется с помощью JWT-токенов."
    )

    doc.add_heading("5. Реализация клиентской части", level=1)
    doc.add_paragraph(
        "Клиентская часть представляет собой одностраничное приложение. Основные страницы: "
        "авторизация и регистрация, список поездок, создание/редактирование поездки, детальная страница поездки "
        "с вкладками маршрута, расходов, чемодана, бронирований, карты и рекомендаций, а также страница настроек."
    )
    doc.add_paragraph(
        "Реализована поддержка тёмной и светлой темы, переключение языка (русский/английский), "
        "выбор валюты, единиц расстояния и температуры."
    )

    doc.add_heading("6. Тестирование", level=1)
    doc.add_paragraph(
        "Проведено функциональное тестирование основных сценариев: регистрация и вход пользователя, "
        "создание поездки, добавление маршрута, расходов и вещей, проверка бюджета, развёртывание в Docker. "
        "Приложение развёрнуто на удалённом сервере и доступно по IP-адресу."
    )

    doc.add_heading("7. Заключение", level=1)
    doc.add_paragraph(
        "В ходе выполнения курсовой работы было разработано полноценное веб-приложение для планирования поездок. "
        "Приложение имеет клиентскую и серверную части, систему аутентификации, возможность развёртывания в Docker. "
        "Полученный результат соответствует поставленным требованиям."
    )

    doc.add_heading("Список использованных источников", level=1)
    sources = [
        "Django Documentation. – https://docs.djangoproject.com/",
        "Django REST Framework Documentation. – https://www.django-rest-framework.org/",
        "React Documentation. – https://react.dev/",
        "TypeScript Documentation. – https://www.typescriptlang.org/",
        "Tailwind CSS Documentation. – https://tailwindcss.com/",
        "Docker Documentation. – https://docs.docker.com/",
    ]
    for i, source in enumerate(sources, 1):
        p = doc.add_paragraph(f"{i}. {source}")
        p.paragraph_format.left_indent = Cm(1)

    doc.save(filename)
    print(f"Создан файл: {filename}")


def create_explanatory_note_pdf(filename):
    pdf = FPDF()
    pdf.add_page()

    # Подключаем шрифт с поддержкой кириллицы
    font_path = r"C:\Windows\Fonts\arial.ttf"
    if not os.path.exists(font_path):
        font_path = r"C:\Windows\Fonts\times.ttf"
    pdf.add_font('Arial', '', font_path, uni=True)
    pdf.add_font('Arial', 'B', font_path, uni=True)

    pdf.set_font('Arial', 'B', 16)
    pdf.cell(0, 10, 'ПОЯСНИТЕЛЬНАЯ ЗАПИСКА', ln=True, align='C')
    pdf.set_font('Arial', '', 12)
    pdf.cell(0, 10, 'к курсовой работе', ln=True, align='C')
    pdf.cell(0, 10, 'Тема: «Разработка веб-приложения для планирования поездок»', ln=True, align='C')
    pdf.ln(10)

    pdf.set_font('Arial', 'B', 14)
    pdf.cell(0, 10, '1. Введение', ln=True)
    pdf.set_font('Arial', '', 12)
    pdf.multi_cell(0, 6, 'В современном мире путешествия стали неотъемлемой частью жизни. '
                          'Для эффективной организации поездок необходимо учитывать множество факторов: '
                          'маршрут, расходы, даты, список необходимых вещей. '
                          'Веб-приложение для планирования поездок позволяет систематизировать весь процесс.')
    pdf.ln(5)

    pdf.set_font('Arial', 'B', 14)
    pdf.cell(0, 10, '2. Технологический стек', ln=True)
    pdf.set_font('Arial', '', 12)
    pdf.multi_cell(0, 6, 'Backend: Python 3.10, Django 5, Django REST Framework, JWT.\n'
                          'Frontend: React 18, TypeScript, Vite, Tailwind CSS.\n'
                          'База данных: SQLite.\n'
                          'Контейнеризация: Docker, Docker Compose, Nginx.\n'
                          'Контроль версий: Git, GitHub.')
    pdf.ln(5)

    pdf.set_font('Arial', 'B', 14)
    pdf.cell(0, 10, '3. Реализация', ln=True)
    pdf.set_font('Arial', '', 12)
    pdf.multi_cell(0, 6, 'Серверная часть включает приложения accounts и trips. '
                          'Клиентская часть представляет собой одностраничное приложение с возможностью '
                          'составления маршрута, учёта расходов, формирования списка вещей и настройками.')
    pdf.ln(5)

    pdf.set_font('Arial', 'B', 14)
    pdf.cell(0, 10, '4. Заключение', ln=True)
    pdf.set_font('Arial', '', 12)
    pdf.multi_cell(0, 6, 'В ходе выполнения курсовой работы было разработано полноценное веб-приложение '
                          'для планирования поездок. Приложение имеет клиентскую и серверную части, '
                          'систему аутентификации и возможность развёртывания в Docker.')

    pdf.output(filename)
    print(f"Создан файл: {filename}")


def create_review_docx(filename):
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(1.5)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("ОТЗЫВ\n")
    run.bold = True
    run.font.size = Pt(16)
    run = p.add_run("на курсовую работу\n")
    run.font.size = Pt(14)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run("Студент: _______________________________\n\n")
    run.font.size = Pt(12)
    run = p.add_run("Группа: ________________________________\n\n")
    run.font.size = Pt(12)
    run = p.add_run("Тема работы: «Разработка веб-приложения для планирования поездок»\n\n")
    run.font.size = Pt(12)

    p = doc.add_paragraph()
    run = p.add_run("Общая характеристика работы:\n\n").bold = True
    doc.add_paragraph(
        "Курсовая работа посвящена разработке веб-приложения для планирования поездок. "
        "В работе реализованы все заявленные функции: создание поездок, составление маршрута, "
        "учёт расходов, список вещей в чемодане, бронирования, настройки пользователя."
    )

    p = doc.add_paragraph()
    run = p.add_run("Достоинства работы:\n").bold = True
    merits = [
        "Полнота реализации заявленного функционала",
        "Использование современных технологий (Django, React, TypeScript, Docker)",
        "Наличие системы аутентификации и настроек пользователя",
        "Качественное оформление пользовательского интерфейса",
        "Возможность развёртывания в Docker-контейнерах",
        "Хорошая структурированность кода и наличие документации",
    ]
    for item in merits:
        doc.add_paragraph(item, style='List Bullet')

    p = doc.add_paragraph()
    run = p.add_run("\nЗамечания:\n").bold = True
    doc.add_paragraph(
        "Рекомендуется добавить автоматические тесты backend-части и расширить интеграцию "
        "с внешними сервисами (Google Maps, Google Calendar) после получения API-ключей."
    )

    p = doc.add_paragraph()
    run = p.add_run("\nЗаключение:\n").bold = True
    doc.add_paragraph(
        "Курсовая работа выполнена на высоком уровне, соответствует требованиям и заслуживает оценки «отлично»."
    )

    p = doc.add_paragraph("\n\n\n")
    run = p.add_run("Руководитель: _______________\n\n")
    run.font.size = Pt(12)
    run = p.add_run("Дата: _______________")
    run.font.size = Pt(12)

    doc.save(filename)
    print(f"Создан файл: {filename}")


if __name__ == '__main__':
    base_dir = os.path.dirname(os.path.abspath(__file__))

    create_explanatory_note_docx(os.path.join(base_dir, 'Пояснительная записка.docx'))
    create_explanatory_note_pdf(os.path.join(base_dir, 'Пояснительная записка.pdf'))
    create_review_docx(os.path.join(base_dir, 'Отзыв.docx'))

    print("\nВсе документы созданы успешно.")
